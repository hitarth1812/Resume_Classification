import asyncio
import io
import unittest

import pymupdf
from docx import Document
from fastapi import HTTPException, UploadFile

from model_api import main


SAMPLE_RESUME = (
    "Experienced accountant with strong skills in financial reporting, bookkeeping, "
    "accounts payable, accounts receivable, payroll, tax preparation, auditing, "
    "budgeting, and Microsoft Excel. Managed monthly close activities and "
    "reconciled general ledger accounts."
)


class ClassifierTests(unittest.TestCase):
    def test_model_returns_configured_category_and_normalized_scores(self):
        category, score, alternatives = main.model_scores(SAMPLE_RESUME)
        self.assertIn(category, main.CONFIG["labels"])
        self.assertGreaterEqual(score, 0.0)
        self.assertLessEqual(score, 1.0)
        self.assertEqual(len(alternatives), 3)
        self.assertTrue(all(item["category"] in main.CONFIG["labels"] for item in alternatives))
        self.assertGreaterEqual(score, max(item["score"] for item in alternatives))
        self.assertLessEqual(score + sum(item["score"] for item in alternatives), 1.0)

    def test_classifier_label_mapping_uses_category_names(self):
        category, _, _ = main.model_scores(SAMPLE_RESUME)
        self.assertNotRegex(category, r"^\d+$")

    def test_text_input_classifies(self):
        result = asyncio.run(main.classify(file=None, text=SAMPLE_RESUME))
        self.assertIn(result["category"], main.CONFIG["labels"])
        self.assertIn("score", result)

    def test_too_short_text_returns_actionable_400(self):
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(main.classify(file=None, text="Too short"))
        self.assertEqual(raised.exception.status_code, 400)

    def test_oversized_text_returns_413(self):
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(main.classify(file=None, text="x" * (main.MAX_TEXT_CHARS + 1)))
        self.assertEqual(raised.exception.status_code, 413)

    def test_file_and_text_together_returns_400(self):
        upload = UploadFile(filename="resume.txt", file=io.BytesIO(SAMPLE_RESUME.encode()))
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(main.classify(file=upload, text=SAMPLE_RESUME))
        self.assertEqual(raised.exception.status_code, 400)

    def test_unsupported_file_returns_415(self):
        upload = UploadFile(filename="resume.exe", file=io.BytesIO(b"not a resume"))
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(main.classify(file=upload, text=""))
        self.assertEqual(raised.exception.status_code, 415)

    def test_oversized_file_returns_413(self):
        upload = UploadFile(filename="resume.txt", file=io.BytesIO(b"x" * (main.MAX_FILE_BYTES + 1)))
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(main.classify(file=upload, text=""))
        self.assertEqual(raised.exception.status_code, 413)

    def test_corrupt_pdf_returns_422_instead_of_server_error(self):
        upload = UploadFile(filename="resume.pdf", file=io.BytesIO(b"not a valid PDF"))
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(main.classify(file=upload, text=""))
        self.assertEqual(raised.exception.status_code, 422)

    def test_pdf_with_embedded_text_classifies(self):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), SAMPLE_RESUME)
        content = pdf.tobytes()
        pdf.close()
        upload = UploadFile(filename="resume.pdf", file=io.BytesIO(content))
        result = asyncio.run(main.classify(file=upload, text=""))
        self.assertIn(result["category"], main.CONFIG["labels"])

    def test_docx_table_text_is_extracted(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = SAMPLE_RESUME
        content = io.BytesIO()
        doc.save(content)
        upload = UploadFile(filename="resume.docx", file=io.BytesIO(content.getvalue()))
        result = asyncio.run(main.classify(file=upload, text=""))
        self.assertIn(result["category"], main.CONFIG["labels"])


if __name__ == "__main__":
    unittest.main()
